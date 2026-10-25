"""PPTX 교안에서 프롬프트를 뽑아 prompts.js를 만든다.

사용법: python3 tools/build_data.py 교안.pptx prompts.js

슬라이드 종류
- 표지(1번)·목차: 건너뛴다
- 챕터 표지(텍스트 1개, 또는 '학습목표'가 있는 [실전] 표지): 새 챕터를 연다
- 프롬프트 장표('핵심' 상자가 있는 장표): 현재 챕터에 항목으로 넣는다
"""
import json
import re
import sys

from pptx import Presentation

# 사용자가 채우는 칸 = (슬라이드 번호, 줄 이름, 대괄호 원문)
# 이 목록과 FILL_ROWS 밖의 대괄호([확인 필요], [추정], [ ]로 비운다 등)는
# AI가 답에 붙이는 표시이므로 노란 글씨로만 보여 준다.
FILL_IN = {
    (39, "요청", "[주제]"), (39, "요청", "[임원/직책자]"), (39, "요청", "[ ]"),
    (51, "기준", "[글로벌 선도기업 / 정부 정책 / 연구기관·컨설팅]"),
    (83, "역할", "[타 부서 임원]"),
    (86, "기준", "[ ]"),
    (116, "요청", "[각 부서 담당자]"), (116, "요청", "[12.15(목) 17:00]"),
    (118, "요청", "[원장님]"), (118, "요청", "[중간보고]"), (118, "요청", "[일시]"),
}
# 이 줄의 대괄호는 모두 사용자가 채우는 칸
FILL_ROWS = {"자료", "나", "독자", "주제", "결정", "방식"}

LABEL_TOP = 1879601   # '핵심' / '프롬프트' 라벨 상자의 위치
MID = 3000000         # 왼쪽(핵심) 열과 오른쪽(프롬프트) 열의 경계


def walk(shapes):
    for s in shapes:
        if s.shape_type == 6:
            yield from walk(s.shapes)
        else:
            yield s


def clean(t):
    return re.sub(r"[ \t]{2,}", " ", " ".join(x.strip() for x in re.split(r"[\n\v]", t))).strip()


def join_wrapped(lines):
    """슬라이드에서 줄 맞춤용으로 끊은 줄을 이어 붙인다. '예:'·번호 줄은 들여쓴 채 둔다."""
    out = []
    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue
        if out and line[:1].isspace():
            body = line.strip()
            if re.match(r"^(예:|예\)|[①-⑩]|출처:)", body):
                out.append("  " + body)
            else:
                out[-1] = out[-1].rstrip() + " " + body
        else:
            out.append(line)
    return out


def segments(slide, label, line):
    segs, pos = [], 0
    for m in re.finditer(r"\[[^\]]*\]", line):
        if label in FILL_ROWS or (slide, label, m.group()) in FILL_IN:
            if m.start() > pos:
                segs.append(line[pos:m.start()])
            segs.append({"f": m.group()[1:-1].strip()})
            pos = m.end()
    if pos < len(line):
        segs.append(line[pos:])
    return segs


def parse_prompt(no, ts):
    near = lambda s, top: abs(s.top - top) < 5000
    core_lbl = next(s for s in ts if s.text_frame.text.strip() == "핵심")
    lab = next(s for s in ts if near(s, LABEL_TOP) and s.left > MID)
    body = max((s for s in ts if s.left > MID and "\t" in s.text_frame.text), key=lambda s: len(s.text_frame.text))
    left = sorted((s for s in ts if s.left < MID and s.top > core_lbl.top + 100000 and s.top < 6000000 and s is not core_lbl),
                  key=lambda s: s.top)
    core, pts = left[0], left[1:]
    item = {"id": no, "label": lab.text_frame.text.strip(), "core": clean(core.text_frame.text), "points": [], "rows": []}
    for i in range(0, len(pts) - 1, 2):
        item["points"].append([clean(pts[i].text_frame.text), clean(pts[i + 1].text_frame.text)])
    for para in body.text_frame.paragraphs:
        if not para.text.strip():
            continue
        if "\t" not in para.text:  # '▶ 다음 장 2/2까지 이어 붙여…' 같은 안내 줄
            item["cont"] = clean(para.text).lstrip("▶ ").strip()
            continue
        name, _, rest = para.text.partition("\t")
        lines = join_wrapped(rest.split("\v"))
        item["rows"].append({"k": name.strip(), "v": [segments(no, name.strip(), l) for l in lines]})
    used = {id(core_lbl), id(lab), id(body), *map(id, left)}
    for s in ts:
        t = s.text_frame.text.strip()
        if id(s) in used or t == "!" or s.top < 100000:
            continue
        if near(s, 215900):
            item["title"] = clean(t)
        elif 1000000 < s.top < 1150000 and s.left < 10000:
            item["sub"] = clean(t)
        elif s.top > 6000000:
            item["tip"] = clean(t)
        else:
            item.setdefault("notes", []).append(t.replace("\v", "\n"))
    return item


def merge_pair(first, second):
    """'프롬프트 1/2'와 '2/2' 장표를 프롬프트 하나로 합친다."""
    first["label"] = "프롬프트"
    first.pop("cont", None)
    # 2/2를 가리키던 말을 합친 프롬프트 안의 항목 이름으로 바꾼다
    target = second["rows"][0]["k"] if second["rows"] else ""
    for row in first["rows"]:
        for line in row["v"]:
            for i, seg in enumerate(line):
                if isinstance(seg, str) and "아래 2/2대로" in seg:
                    line[i] = seg.replace("아래 2/2대로", f"아래 '{target}' 항목대로")
    first["rows"] += second["rows"]
    # 2/2의 '사용법' 중 이어 붙이기 안내는 빼고 나머지를 이어 번호를 매긴다
    extra = [p for p in second["points"] if "이어 붙이기" not in p[0]]
    n = len(first["points"])
    for i, (head, desc) in enumerate(extra, n + 1):
        first["points"].append([re.sub(r"^\d+", f"{i:02d}", head), desc])
    if second.get("tip") and not first.get("tip"):
        first["tip"] = second["tip"]


def main(src, dst):
    prs = Presentation(src)
    chapters, cur = [], None
    for no, slide in enumerate(prs.slides, 1):
        ts = [s for s in walk(slide.shapes) if s.has_text_frame and s.text_frame.text.strip()]
        texts = [s.text_frame.text.strip() for s in ts]
        if no == 1 or "목차" in texts:
            continue
        if len(ts) == 1:  # 챕터 표지
            cur = {"title": texts[0], "items": []}
            chapters.append(cur)
            continue
        if "학습목표" in texts:  # [실전] 표지
            title = next(t for t in texts if t.startswith("[실전"))
            goals = [g.lstrip("▶ ").strip() for t in texts if t.startswith("▶") for g in re.split(r"\n", t) if g.strip()]
            goals = [re.sub(r"\s{2,}", " ", g.lstrip("▶").strip()) for g in goals]
            cur = {"title": title, "goals": goals, "items": []}
            chapters.append(cur)
            continue
        if "핵심" not in texts:
            print(f"건너뜀: {no}번 슬라이드", file=sys.stderr)
            continue
        item = parse_prompt(no, ts)
        prev = cur["items"][-1] if cur["items"] else None
        if item["label"].endswith("2/2") and prev and prev["label"].endswith("1/2"):
            merge_pair(prev, item)
            continue
        cur["items"].append(item)

    with open(dst, "w", encoding="utf-8") as f:
        f.write("// tools/build_data.py로 생성한 파일입니다.\n")
        f.write("window.PROMPT_DATA = [\n")
        for ci, ch in enumerate(chapters):
            head = {k: v for k, v in ch.items() if k != "items"}
            f.write(json.dumps(head, ensure_ascii=False)[:-1] + ', "items": [\n')
            f.write(",\n".join(json.dumps(it, ensure_ascii=False) for it in ch["items"]))
            f.write("\n]}" + ("," if ci < len(chapters) - 1 else "") + "\n")
        f.write("];\n")
    print(f"{len(chapters)}개 챕터, {sum(len(c['items']) for c in chapters)}개 항목")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
